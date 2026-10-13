# dianyicha_price.py
"""Jiangsu day-ahead prices, read from the 电易电查 workbook.

`docs/日前价格261009.xlsx` carries four 96-point curves and nothing else:

| sheet | what its own column header calls it |
|---|---|
| `261009` | 日前均值（CNY/MWh） |
| `7days` | 日前均价 |
| `30days` | 日前均价 |
| `halfyear` | 日前均价 |

**All four derive from real historical market data, and all four are spatially
averaged.** Two questions are kept apart, because this file answers them
differently and a single flag cannot carry both:

* **`spatial_statistic`** -- what was averaged over *space*. `spatial_mean` for
  every curve here: each 96-point series is an average across locations, so no
  individual bus, zone or settlement point is represented by it. The spatial
  extent, the weighting and the number of points behind it are **not recorded in
  the workbook** and are left unverified rather than assumed.
* **`temporal_statistic`** -- what was averaged over *time*. `261009` is a
  **single day**; `7days`, `30days` and `halfyear` are means over a window of
  7, 30 and 180 days.

**`temporal_statistic` is the sole authority for single-day identity.** Nothing
downstream decides that question from `role`, from `is_measured`, or from the
presence of 均值 in a header.

**`261009`'s single-day reading is a project decision, not an external
confirmation.** Its column header says `日前均值`, "day-ahead mean", which at
face value reads as a temporal mean. It is read here as a *spatial* mean, and
the date in the sheet name (26/10/09) as the day the curve covers. The workbook
cannot settle this -- it was assembled by hand and carries no platform export
metadata -- so `CURVE_META` records the ruling, who made it and when, and
`externally_confirmed` stays False until the provider's own column definition is
obtained.

What survives either reading is the quantity a training split depends on: this
file holds **one** real day and three window means, so it is not a *series* of
daily prices. `contains_multi_day_series` is False throughout, and a study built
on this file cannot say anything about day-to-day variation or about
generalisation across days.

**Deprecated compatibility fields.** `is_measured`, `aggregation`,
`mean_window_days` and `describe()["contains_measured_daily_series"]` are
retained for one compatibility window and are **derived properties**, not
independent facts. Their meanings are listed in `DEPRECATED_FIELDS` and
published by `describe()` under `deprecated`. `is_measured` in particular no
longer means "confirmed single day" -- it is True for all four curves, because
all four are real measured data. Anything that still reads it for single-day
identity reads the opposite of what it intends.

**Time labels are interval-END, and the mapping is positional.** A label is the
*close* of the interval it belongs to: `00:15` closes the interval that runs
`00:00`-`00:15`, and `24:00` closes `23:45`-`24:00`. The simulator's period `t`
is `[t*15min, (t+1)*15min)`, so label `i` belongs to period `i` and the mapping
is the identity. Nothing is shifted.

That is asserted rather than assumed. The expected sequence is
`00:15, 00:30, ..., 24:00`, i.e. `15 * (i + 1)` minutes past midnight for slot
`i`, and a file that disagrees raises. A workbook using interval-START labels
(`00:00`-`23:45`) is therefore rejected with a message naming the convention
instead of being silently rotated by one period — a one-period rotation is
exactly the kind of error that survives review, because every number still looks
plausible.

The last label needs care: `24:00` does not exist as a `datetime.time`, so Excel
stores it as a `datetime` at midnight of its epoch date, `1900-01-01 00:00:00`.
It is accepted **only** in the final position; a midnight `datetime` anywhere
else is an error rather than a guess.

The workbook is read with `openpyxl` in read-only, values-only mode and is never
written to. No text or CSV copy is produced: this module is the reader, and the
provider in `wholesale_price.py` calls it directly.
"""

from __future__ import annotations

import argparse
import datetime as _dt
import hashlib
import json
import sys
from dataclasses import dataclass, field, replace
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

import numpy as np

#: Periods in a day, matching the simulator's horizon.
N_SLOTS = 96
SLOT_MINUTES = 15

#: Sheets, in the order a scenario enumerates them. The order is the `day_id`
#: mapping the provider uses, so it is part of the interface, not a convention.
SHEETS: Tuple[str, ...] = ("261009", "7days", "30days", "halfyear")

#: Declared vocabularies for the two statistics. Kept closed so a typo cannot
#: invent a third category that no consumer knows how to read.
SPATIAL_STATISTICS: Tuple[str, ...] = ("spatial_mean", "unspecified")
TEMPORAL_STATISTICS: Tuple[str, ...] = ("single_day", "window_mean",
                                        "unspecified")

#: Source kinds. `real_market_data` means the numbers are derived from real
#: measured market data, as opposed to a simulation. It says nothing about
#: aggregation.
SOURCE_KINDS: Tuple[str, ...] = ("real_market_data", "unknown")

#: Human-readable label per sheet. **Descriptive only** -- `role` carries no
#: authority over what a curve is; `temporal_statistic` does. It is kept
#: because it reads well in a table, and `261009`'s label was corrected from
#: `same_day_or_same_day_mean` once the ambiguity it named was resolved.
ROLES: Dict[str, str] = {
    "261009": "single_day_spatial_mean",
    "7days": "trailing_7_day_mean",
    "30days": "trailing_30_day_mean",
    "halfyear": "trailing_180_day_mean",
}

#: Fields kept only for a compatibility window. Each entry states, in the
#: published payload, what the field now means -- because every one of them has
#: changed meaning or been superseded, and a downstream caller still reading the
#: old sense would get a confident wrong answer rather than an error.
#:
#: `describe()` publishes this mapping under `deprecated`, and a test asserts the
#: deprecated fields are present and carry these semantics, so the old meanings
#: cannot be quietly reintroduced.
DEPRECATED_FIELDS: Dict[str, str] = {
    "is_measured": (
        "True when the curve's underlying source is real market measured data. "
        "It does NOT mean the curve is a single day, and it does NOT mean it is "
        "free of spatial averaging -- it is True for all four curves here. Use "
        "`temporal_statistic` for single-day identity."),
    "aggregation": (
        "Legacy single-axis label, superseded by `temporal_statistic`. It could "
        "not express that a curve is both spatially averaged and a single day, "
        "which is why `261009` no longer reads `unresolved`."),
    "mean_window_days": (
        "Legacy name for `window_days` that also encodes the temporal "
        "statistic: it is None unless the curve is a window mean. Read "
        "`window_days` with `temporal_statistic` instead."),
    "contains_measured_daily_series": (
        "Whether the set holds a per-day sample series spanning MULTIPLE real "
        "dates. False here, and unaffected by `261009` -- one day is not a "
        "series. It does NOT say whether a real single-day 96-period curve is "
        "present; use `contains_single_day_curve` for that."),
    "contains_confirmed_daily_series": (
        "Whether any curve's SINGLE-DAY reading is confirmed from outside the "
        "project. False here: 261009 is read as a single day by project ruling "
        "and the workbook cannot confirm it. It does NOT mean no single-day "
        "curve is present; use `contains_single_day_curve` for that, and "
        "`externally_confirmed` per curve for the confirmation itself."),
}

#: Legacy `aggregation` values, derived from `temporal_statistic` so the field
#: cannot drift from the authority beside it.
_LEGACY_AGGREGATION: Dict[str, str] = {
    "single_day": "single_day",
    "window_mean": "trailing_mean",
    "unspecified": "unresolved",
}

#: Per-curve provenance, as far as it can be established.
#:
#: Four questions, answered separately: **where the numbers came from**
#: (`source_kind`), **what was averaged over space** (`spatial_statistic`),
#: **what was averaged over time** (`temporal_statistic`), and **over how many
#: days** (`window_days`). They are four fields rather than one because one
#: cannot carry them: a curve is routinely both spatially averaged *and* a
#: single day, and the retired `aggregation` axis had no way to say so -- which
#: is exactly why `261009` was mislabelled `unresolved`.
#:
#: **All four curves are `spatial_mean` over `real_market_data`.** The spatial
#: extent, the weighting scheme and the number of underlying points are **not
#: recorded in the workbook**; they are left `None` rather than guessed.
#:
#: **`261009` is recorded as `single_day` by project ruling, not by external
#: confirmation.** Its header reads `日前均值` ("day-ahead mean"), which at face
#: value is a *temporal* mean; the ruling reads that 均值 as a *spatial* one and
#: the sheet's date 26/10/09 as the day covered. Nothing in the workbook settles
#: it: the file was hand-assembled (`docProps` record a personal author and a
#: creation date well before the sheet's own date) and carries no platform
#: export metadata, so there is no column definition to consult here.
#: `externally_confirmed` is therefore False and stays so until the provider's
#: own definition of the column is obtained. `declared_by` names the ruling.
#:
#: `window_start` / `window_end` are `None` for every curve: the workbook does
#: not state the trailing windows' end dates, and inventing them would be worse
#: than leaving them blank. See `reports/final_research_summary.md`.
CURVE_META: Dict[str, Dict[str, object]] = {
    "261009": {
        "source_kind": "real_market_data",
        "spatial_statistic": "spatial_mean",
        "spatial_reference": None,
        "temporal_statistic": "single_day",
        "window_days": 1,
        "reference_period": "26/10/09",
        "externally_confirmed": False,
        "declared_by": "project ruling 2026-10-12",
        "temporal_readings": (
            "read as a single day's curve; the header's 均值 is taken as a "
            "spatial mean rather than a temporal one. The workbook does not "
            "settle this and it awaits the provider's column definition."),
        "window_start": None,
        "window_end": None,
        "header_label": "日前均值（CNY/MWh）",
    },
    "7days": {
        "source_kind": "real_market_data",
        "spatial_statistic": "spatial_mean",
        "spatial_reference": None,
        "temporal_statistic": "window_mean",
        "window_days": 7,
        "reference_period": None,
        "externally_confirmed": True,
        "declared_by": "the workbook's own column header (日前均价)",
        "temporal_readings": "a trailing mean over 7 days of real data",
        "window_start": None,
        "window_end": None,
        "header_label": "日前均价",
    },
    "30days": {
        "source_kind": "real_market_data",
        "spatial_statistic": "spatial_mean",
        "spatial_reference": None,
        "temporal_statistic": "window_mean",
        "window_days": 30,
        "reference_period": None,
        "externally_confirmed": True,
        "declared_by": "the workbook's own column header (日前均价)",
        "temporal_readings": "a trailing mean over 30 days of real data",
        "window_start": None,
        "window_end": None,
        "header_label": "日前均价",
    },
    "halfyear": {
        "source_kind": "real_market_data",
        "spatial_statistic": "spatial_mean",
        "spatial_reference": None,
        "temporal_statistic": "window_mean",
        "window_days": 180,
        "reference_period": None,
        "externally_confirmed": True,
        "declared_by": "the workbook's own column header (日前均价)",
        "temporal_readings": "a trailing mean over half a year of real data",
        "window_start": None,
        "window_end": None,
        "header_label": "日前均价",
    },
}

#: Units, as declared by every sheet's own header.
UNITS = "CNY/MWh"

#: The file name, relative to the repository root.
DEFAULT_WORKBOOK = Path("docs") / "日前价格261009.xlsx"


def resolve_workbook(path: Optional[str] = None) -> Path:
    """Locate the price workbook, defaulting to the one under `docs/`."""
    if path is not None:
        p = Path(path)
        if not p.exists():
            raise FileNotFoundError(f"price workbook not found: {p}")
        return p
    root = Path(__file__).resolve().parent.parent
    candidate = root / DEFAULT_WORKBOOK
    if not candidate.exists():
        raise FileNotFoundError(
            f"price workbook not found: {candidate}. Pass an explicit path.")
    return candidate


@dataclass(frozen=True)
class Curve:
    """One 96-point price curve, with what it is and where it came from.

    Four independent facts, and the reason they are four: a curve is routinely
    both spatially averaged and a single day, so no single axis can carry them.
    `temporal_statistic` alone answers "is this one day"; `role` and the
    deprecated properties below are labels and must not be read as evidence.
    """

    sheet: str
    #: Human-readable label. **Descriptive only** -- derived from nothing and
    #: authoritative over nothing.
    role: str
    prices: Tuple[float, ...]
    #: "real_market_data" for every curve here: all four are derived from real
    #: measured market data, as opposed to simulation.
    source_kind: str = "unknown"
    #: What was averaged over **space**; `spatial_mean` for all four, so no
    #: individual bus or zone is represented by any curve. The spatial extent,
    #: the weighting and the number of points are unverified and not guessed.
    spatial_statistic: str = "unspecified"
    #: What was averaged over **time**. **The sole authority** for single-day
    #: identity: `"single_day"` or `"window_mean"`.
    temporal_statistic: str = "unspecified"
    #: Days behind the curve -- 1 for a single day, the window length for a mean.
    window_days: Optional[int] = None
    #: The date the sheet names, where it names one.
    reference_period: Optional[str] = None
    #: True only where a source states the temporal reading. False where the
    #: project ruled on it, as for `261009`.
    externally_confirmed: bool = False
    #: Who or what settled the temporal reading; the ruling, or the header.
    declared_by: Optional[str] = None
    #: The trailing window's end date, or None when it could not be verified.
    window_end: Optional[str] = None
    provenance: Dict[str, object] = field(default_factory=dict)

    # -- the one predicate consumers should use ------------------------------

    @property
    def is_single_day(self) -> bool:
        """Whether this curve is one day. Reads `temporal_statistic` alone."""
        return self.temporal_statistic == "single_day"

    # -- deprecated derived fields, retained for one compatibility window ----
    #
    # These were fields. They are properties now so they cannot drift from the
    # facts beside them. Each changed meaning or was superseded; see
    # `DEPRECATED_FIELDS` for the published statement of what each one is.

    @property
    def is_measured(self) -> bool:
        """DEPRECATED: True when the source is real measured market data.

        It is **True for all four curves**, including the window means, and says
        nothing about being a single day or about spatial averaging. Reading it
        as "confirmed single day" now yields the opposite of the intent.
        """
        return self.source_kind == "real_market_data"

    @property
    def aggregation(self) -> str:
        """DEPRECATED: legacy single-axis label; use `temporal_statistic`."""
        return _LEGACY_AGGREGATION.get(self.temporal_statistic, "unresolved")

    @property
    def mean_window_days(self) -> Optional[int]:
        """DEPRECATED: `window_days`, but None unless it is a window mean."""
        if self.temporal_statistic != "window_mean":
            return None
        return self.window_days

    def array(self) -> np.ndarray:
        return np.asarray(self.prices, dtype=float)

    def stats(self) -> Dict[str, float]:
        arr = self.array()
        return {
            "min_cny_mwh": float(arr.min()),
            "max_cny_mwh": float(arr.max()),
            "mean_cny_mwh": float(arr.mean()),
            "peak_slot": int(np.argmax(arr)),
            "trough_slot": int(np.argmin(arr)),
        }

    def window(self, t_start: int, T: int) -> np.ndarray:
        """`T` periods beginning at `t_start`, the provider's window contract."""
        arr = self.array()
        if t_start < 0 or T <= 0 or t_start + T > N_SLOTS:
            raise ValueError(
                f"window [{t_start}, {t_start + T}) does not fit a "
                f"{N_SLOTS}-period day")
        return arr[t_start:t_start + T]


# ---------------------------------------------------------------------------
# Reading
# ---------------------------------------------------------------------------


def label_minutes(cell, index: int) -> int:
    """Minutes past midnight for a slot label; `24:00` returns 1440.

    Excel cannot hold `24:00` in a `time` cell, so the closing label of the day
    arrives as midnight of the epoch date. It is only meaningful in the final
    slot; anywhere else it is ambiguous and raises.
    """
    if isinstance(cell, _dt.datetime):
        if (cell.hour, cell.minute, cell.second) != (0, 0, 0):
            raise ValueError(
                f"slot {index}: datetime label {cell!r} is not midnight; the "
                "only datetime this reader accepts is the end-of-day 24:00")
        if index != N_SLOTS - 1:
            raise ValueError(
                f"slot {index}: a midnight datetime label is only the 24:00 "
                "end-of-day sentinel, which belongs in the last slot")
        return 24 * 60
    if isinstance(cell, _dt.time):
        if cell.second or cell.microsecond:
            raise ValueError(
                f"slot {index}: label {cell!r} has sub-minute precision")
        return cell.hour * 60 + cell.minute
    raise ValueError(
        f"slot {index}: label {cell!r} is {type(cell).__name__}, not a time")


def _check_slot_axis(labels: Sequence) -> None:
    """Assert the labels are exactly the interval-END clock, in order.

    The expected value at slot `i` is `15 * (i + 1)` minutes. Accepting
    interval-START labels here would silently rotate every price by one period,
    so a mismatch is an error naming the convention, not a repair.
    """
    if len(labels) != N_SLOTS:
        raise ValueError(
            f"expected {N_SLOTS} slot labels, found {len(labels)}")
    minutes = [label_minutes(cell, i) for i, cell in enumerate(labels)]
    for i, got in enumerate(minutes):
        want = SLOT_MINUTES * (i + 1)
        if got != want:
            raise ValueError(
                f"slot {i}: label is {got} minutes past midnight, expected "
                f"{want} ({want // 60:02d}:{want % 60:02d}). The labels are not "
                "the interval-END clock the simulator's periods are aligned to; "
                "an interval-START file must be converted deliberately rather "
                "than read as if it were this one.")
    if len(set(minutes)) != N_SLOTS:
        raise ValueError("slot labels repeat")


def _check_price(cell, index: int, sheet: str) -> float:
    if cell is None:
        raise ValueError(f"{sheet} slot {index}: price is blank")
    if isinstance(cell, bool) or not isinstance(cell, (int, float)):
        try:
            cell = float(str(cell).strip())
        except (TypeError, ValueError):
            raise ValueError(
                f"{sheet} slot {index}: price {cell!r} is not a number")
    value = float(cell)
    if not np.isfinite(value):
        raise ValueError(f"{sheet} slot {index}: price {cell!r} is not finite")
    return value


def read_curves(path: Optional[str] = None,
                sheets: Sequence[str] = SHEETS) -> Dict[str, Curve]:
    """Read every requested sheet into a `Curve`.

    Raises on a missing sheet, a missing or non-numeric price, a repeated or
    misordered slot, or a label that is not the interval-END clock.
    """
    import openpyxl

    wb_path = resolve_workbook(path)
    digest = hashlib.sha256(wb_path.read_bytes()).hexdigest()
    producer = _producer_properties(wb_path)

    wb = openpyxl.load_workbook(wb_path, read_only=True, data_only=True)
    try:
        missing = [s for s in sheets if s not in wb.sheetnames]
        if missing:
            raise KeyError(
                f"sheet(s) {missing} not in {wb_path.name}; it carries "
                f"{wb.sheetnames}")

        curves: Dict[str, Curve] = {}
        for name in sheets:
            ws = wb[name]
            rows = [r for r in ws.iter_rows(values_only=True)
                    if any(c is not None for c in r)]
            if not rows:
                raise ValueError(f"{name}: sheet is empty")

            header = rows[0]
            body = rows[1:]
            if len(body) != N_SLOTS:
                raise ValueError(
                    f"{name}: expected {N_SLOTS} data rows, found {len(body)}")

            labels = [r[0] for r in body]
            _check_slot_axis(labels)
            prices = tuple(_check_price(r[1] if len(r) > 1 else None, i, name)
                           for i, r in enumerate(body))

            meta = CURVE_META.get(name, {})
            curve = Curve(
                sheet=name,
                role=ROLES.get(name, "unspecified"),
                prices=prices,
                source_kind=str(meta.get("source_kind", "unknown")),
                spatial_statistic=str(
                    meta.get("spatial_statistic", "unspecified")),
                temporal_statistic=str(
                    meta.get("temporal_statistic", "unspecified")),
                window_days=meta.get("window_days"),
                window_end=meta.get("window_end"),
                reference_period=meta.get("reference_period"),
                externally_confirmed=bool(meta.get("externally_confirmed",
                                                   False)),
                declared_by=meta.get("declared_by"),
            )
            # Built from the constructed curve rather than from `meta`, so the
            # deprecated copies are derived from the same place the fields are
            # and cannot disagree with them.
            curves[name] = replace(curve, provenance={
                "workbook": wb_path.name,
                "workbook_path": str(wb_path),
                "workbook_sha256": digest,
                "sheet": name,
                "header_label": str(header[1]) if len(header) > 1 else "",
                "units": UNITS,
                "interval_convention": "label_is_interval_end",
                "slot_mapping": "label_index == period_index",
                "role": curve.role,
                # The four facts, carried through so a consumer does not have
                # to re-read this module's constants.
                "source_kind": curve.source_kind,
                "spatial_statistic": curve.spatial_statistic,
                "spatial_reference": meta.get("spatial_reference"),
                "temporal_statistic": curve.temporal_statistic,
                "window_days": curve.window_days,
                "temporal_readings": meta.get("temporal_readings"),
                "reference_period": curve.reference_period,
                "externally_confirmed": curve.externally_confirmed,
                "declared_by": curve.declared_by,
                "window_start": meta.get("window_start"),
                "window_end": curve.window_end,
                # Deprecated copies, derived. Their stated meaning is published
                # by DEPRECATED_FIELDS; `is_measured_semantics` travels with the
                # value so a reader of the JSON alone cannot take it for a
                # single-day claim.
                "is_measured": curve.is_measured,
                "is_measured_semantics": DEPRECATED_FIELDS["is_measured"],
                "aggregation": curve.aggregation,
                "aggregation_semantics": DEPRECATED_FIELDS["aggregation"],
                "mean_window_days": curve.mean_window_days,
                "mean_window_days_semantics":
                    DEPRECATED_FIELDS["mean_window_days"],
                # The workbook's own document properties, which are the only
                # producer-side metadata the file carries.
                "workbook_author": producer.get("author"),
                "workbook_created": producer.get("created"),
                "workbook_modified": producer.get("modified"),
            })
        return curves
    finally:
        wb.close()


def _producer_properties(path: Path) -> Dict[str, object]:
    """The workbook's own `docProps/core.xml` fields.

    The only producer-side metadata the file carries: a personal author and a
    creation date. Worth recording because it is what establishes the workbook
    as a hand-assembled extraction rather than a platform export -- which is in
    turn why no window dates are present in it.
    """
    import re
    import zipfile

    out: Dict[str, object] = {"author": None, "created": None, "modified": None}
    try:
        with zipfile.ZipFile(path) as z:
            xml = z.read("docProps/core.xml").decode("utf-8", "replace")
    except (KeyError, OSError, zipfile.BadZipFile):
        return out
    for key, tag in (("author", "dc:creator"), ("created", "dcterms:created"),
                     ("modified", "dcterms:modified")):
        m = re.search(rf"<{tag}[^>]*>([^<]*)</{tag}>", xml)
        if m:
            out[key] = m.group(1)
    return out


def describe(curves: Dict[str, Curve]) -> Dict[str, object]:
    """A report-ready description of a curve set.

    Says what each curve IS, not just what it is called. Three facts matter for
    every downstream claim, and they are three because they come apart:

    * **All four derive from real measured market data** (`source_kind`).
    * **All four are spatially averaged** (`spatial_statistic`), so none of them
      represents an individual bus, zone or settlement point. The extent and
      the weighting are unverified.
    * **One is a single day and three are window means** (`temporal_statistic`).

    The third is what a training split turns on, and it is answered by
    `temporal_statistic` alone. `contains_multi_day_series` is the flag that
    states the consequence: one day plus three window means is not a *series* of
    daily prices, so nothing here supports a statement about day-to-day
    variation or generalisation across days.

    `contains_single_day_curve` is True and does **not** contradict that -- a
    single-day curve and a multi-day series are different objects, and the
    retired single flag could not hold both.
    """
    temporals = sorted({c.temporal_statistic for c in curves.values()})
    spatials = sorted({c.spatial_statistic for c in curves.values()})
    return {
        "source": "dianyicha",
        "units": UNITS,
        "n_slots": N_SLOTS,
        "interval_convention": "label_is_interval_end",
        "source_kind": "real_market_data",
        "spatial_statistics_present": spatials,
        "temporal_statistics_present": temporals,
        "contains_single_day_curve": any(c.is_single_day
                                         for c in curves.values()),
        "contains_multi_day_series": False,
        # Deprecated. Derived, so it tracks `externally_confirmed` per curve
        # rather than being pinned to a literal.
        "contains_confirmed_daily_series": any(
            c.is_single_day and c.externally_confirmed
            for c in curves.values()),
        "contains_confirmed_daily_series_semantics":
            DEPRECATED_FIELDS["contains_confirmed_daily_series"],
        # Kept under its original name for one compatibility window. Its legacy
        # meaning is about a series over multiple real dates, NOT about whether
        # a real single-day curve is present; see DEPRECATED_FIELDS.
        "contains_measured_daily_series": False,
        "contains_measured_daily_series_semantics":
            DEPRECATED_FIELDS["contains_measured_daily_series"],
        "deprecated": dict(DEPRECATED_FIELDS),
        "note": (
            "All four curves derive from real measured market data and all four "
            "are SPATIAL means -- no bus or zone is represented by any of them, "
            "and the spatial extent and weighting are not recorded in the "
            "workbook. They differ in their TEMPORAL statistic: 261009 is a "
            "single day and 7days, 30days and halfyear are means over 7, 30 and "
            "180 days. 261009's single-day reading is a PROJECT RULING reading "
            "its header's 均值 as a spatial mean, not an external confirmation. "
            "One day plus three window means is not a series of daily prices, "
            "so results built on these curves are price-scenario comparisons, "
            "not evidence about day-to-day behaviour, volatility or "
            "generalisation across days."),
        "curves": {
            name: {
                "role": c.role,
                "source_kind": c.source_kind,
                "spatial_statistic": c.spatial_statistic,
                "temporal_statistic": c.temporal_statistic,
                "window_days": c.window_days,
                "is_single_day": c.is_single_day,
                "externally_confirmed": c.externally_confirmed,
                "declared_by": c.declared_by,
                "reference_period": c.reference_period,
                "window_end": c.window_end,
                "header_label": c.provenance.get("header_label"),
                # Deprecated, derived; not authoritative for anything above.
                "is_measured": c.is_measured,
                "aggregation": c.aggregation,
                "mean_window_days": c.mean_window_days,
                **c.stats(),
            }
            for name, c in curves.items()
        },
    }


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        description="Read the 电易电查 day-ahead price workbook.")
    p.add_argument("--workbook", default=None, help="path to the workbook")
    p.add_argument("--sheet", action="append", dest="sheets", default=None,
                   help="sheet to read (repeatable); default is all four")
    p.add_argument("--json", action="store_true", help="machine-readable output")
    return p


def main(argv: Optional[Sequence[str]] = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        curves = read_curves(args.workbook,
                             tuple(args.sheets) if args.sheets else SHEETS)
    except Exception as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2

    payload = describe(curves)
    if args.json:
        print(json.dumps(payload, ensure_ascii=False, indent=2))
        return 0

    print(f"source: dianyicha   units: {UNITS}   slots: {N_SLOTS}")
    print(payload["note"])
    print()
    print(f"{'sheet':<10}{'temporal':<12} {'days':>5} {'spatial':<13} "
          f"{'role':<24}{'min':>9}{'max':>9}{'mean':>9}{'peak':>6}{'trough':>7}"
          f"  header")
    for name, info in payload["curves"].items():
        print(f"{name:<10}{info['temporal_statistic']:<12} "
              f"{info['window_days']:>5} {info['spatial_statistic']:<13} "
              f"{info['role']:<24}{info['min_cny_mwh']:>9.2f}"
              f"{info['max_cny_mwh']:>9.2f}{info['mean_cny_mwh']:>9.2f}"
              f"{info['peak_slot']:>6}{info['trough_slot']:>7}"
              f"  {info['header_label']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
